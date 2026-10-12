"""
Attendance Marker — Full System Credential Rotation & Live Verification.
Rotates all 128 active user accounts:
- 119 Students
- 8 Faculty
- 1 Admin

Generates unique, cryptographically random, policy-compliant 14-char passwords.
Hashes via Bcrypt, updates database state, commits, and performs live API login verification.
Outputs credentials strictly to ATTENDANCE_MARKER_ALL_ACCOUNT_CREDENTIALS.xlsx.
NO plaintext passwords are printed to terminal, audit logs, or reports.
"""
import os
import sys
import json
import secrets
import string
import datetime
from datetime import timezone
import urllib.request
import urllib.error

sys.path.insert(0, os.path.abspath("."))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.database import SessionLocal
from app.models.user import User, UserRole, Faculty
from app.models.student import Student
from app.models.audit import AuditLog
from app.core.security import get_password_hash

BASE_URL = "http://127.0.0.1:8000"
EXCEL_OUTPUT_PATH = r"C:\Users\Abishek\Desktop\am-oss\ATTENDANCE_MARKER_ALL_ACCOUNT_CREDENTIALS.xlsx"

def generate_strong_password(length: int = 14) -> str:
    safe_symbols = "!@#$%^&*"
    chars = string.ascii_letters + string.digits + safe_symbols
    pwd = [
        secrets.choice(string.ascii_uppercase),
        secrets.choice(string.ascii_lowercase),
        secrets.choice(string.digits),
        secrets.choice(safe_symbols),
        secrets.choice(string.ascii_uppercase),
        secrets.choice(string.ascii_lowercase),
        secrets.choice(string.digits),
        secrets.choice(safe_symbols),
    ]
    pwd += [secrets.choice(chars) for _ in range(length - len(pwd))]
    secrets.SystemRandom().shuffle(pwd)
    return "".join(pwd)

def http_post(endpoint: str, payload: dict):
    url = f"{BASE_URL}{endpoint}"
    data_bytes = json.dumps(payload).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    req = urllib.request.Request(url, data=data_bytes, headers=headers, method="POST")
    try:
        with urllib.request.urlopen(req) as resp:
            body = resp.read().decode("utf-8")
            return resp.status, json.loads(body) if body else {}
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8")
        try:
            parsed = json.loads(body)
        except Exception:
            parsed = {"detail": body}
        return e.code, parsed

def execute_rotation():
    print("=" * 75)
    print("STARTING FULL SYSTEM CREDENTIAL ROTATION (128 ACTIVE ACCOUNTS)")
    print("=" * 75)

    db = SessionLocal()

    # STEP 1: INVENTORY ALL ACCOUNTS
    print("\n[STEP 1] Inventorying all active accounts from attendance_system.db...")
    active_users = db.query(User).filter(User.is_active == True).order_by(User.id).all()
    total_count = len(active_users)
    print(f"  Total active users found: {total_count}")
    if total_count != 128:
        raise ValueError(f"Expected exactly 128 active accounts, found {total_count}! Aborting.")

    admin_users = [u for u in active_users if u.role == UserRole.ADMIN]
    faculty_users = [u for u in active_users if u.role == UserRole.FACULTY]
    student_users = [u for u in active_users if u.role == UserRole.STUDENT]

    print(f"  - Admin Accounts:   {len(admin_users)} (Expected: 1)")
    print(f"  - Faculty Accounts: {len(faculty_users)} (Expected: 8)")
    print(f"  - Student Accounts: {len(student_users)} (Expected: 119)")

    assert len(admin_users) == 1, f"Expected 1 admin, got {len(admin_users)}"
    assert len(faculty_users) == 8, f"Expected 8 faculty, got {len(faculty_users)}"
    assert len(student_users) == 119, f"Expected 119 students, got {len(student_users)}"

    # STEP 2: GENERATE NEW PASSWORDS
    print("\n[STEP 2] Generating unique cryptographically random passwords...")
    generated_passwords = set()
    account_records = []

    generation_time_str = datetime.datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    for user in active_users:
        while True:
            pwd = generate_strong_password(14)
            if pwd not in generated_passwords:
                generated_passwords.add(pwd)
                break

        # Collect metadata
        full_name = user.username
        roll_number = ""
        employee_id = ""

        if user.role == UserRole.STUDENT and user.student_profile:
            full_name = user.student_profile.full_name
            roll_number = user.student_profile.roll_number
        elif user.role == UserRole.FACULTY and user.faculty_profile:
            full_name = user.faculty_profile.full_name
            employee_id = user.faculty_profile.employee_id
        elif user.role == UserRole.ADMIN:
            full_name = "Administrator"

        account_records.append({
            "user_id": user.id,
            "role": user.role.value,
            "full_name": full_name,
            "username": user.username,
            "email": user.email or "",
            "roll_number": roll_number,
            "employee_id": employee_id,
            "new_password": pwd,
            "force_password_change": True,
            "user_obj": user,
            "verified_login": False,
            "generated_at": generation_time_str
        })

    print(f"  Generated {len(generated_passwords)} distinct passwords for 128 accounts.")
    assert len(generated_passwords) == 128, "Passwords are not 100% unique!"

    # STEP 3: RESET PASSWORDS IN DATABASE
    print("\n[STEP 3] Hashing passwords and updating database in atomic transaction...")
    now_utc = datetime.datetime.now(timezone.utc)

    for item in account_records:
        user = item["user_obj"]
        user.hashed_password = get_password_hash(item["new_password"])
        user.force_password_change = True
        user.failed_login_attempts = 0
        user.locked_until = None
        user.session_token = None

    # Insert audit log
    admin_id = admin_users[0].id
    audit = AuditLog(
        user_id=admin_id,
        action="SYSTEM_WIDE_CREDENTIAL_ROTATION",
        target="ALL_ACTIVE_ACCOUNTS",
        details="Rotated credentials for 128 accounts (119 students, 8 faculty, 1 admin) with forced password change.",
        timestamp=now_utc
    )
    db.add(audit)
    db.commit()
    print("  Database updated and committed successfully.")

    # STEP 4: VERIFY EACH PASSWORD AGAINST LIVE BACKEND
    print("\n[STEP 4] Verifying all 128 passwords against live backend endpoints...")
    verification_failures = []
    verified_count = 0

    for idx, item in enumerate(account_records, 1):
        role = item["role"]
        pwd = item["new_password"]
        user_id = item["user_id"]

        if role == "STUDENT":
            endpoint = "/api/v1/auth/student/login"
            payload = {"identifier": item["roll_number"], "password": pwd}
        elif role == "FACULTY":
            endpoint = "/api/v1/auth/login"
            payload = {"username_or_email": item["employee_id"], "password": pwd}
        else:  # ADMIN
            endpoint = "/api/v1/auth/login"
            payload = {"username_or_email": item["username"], "password": pwd}

        status, resp_data = http_post(endpoint, payload)

        if status != 200:
            err_msg = f"Account {role} ID {user_id} ({item['username']}) failed authentication: HTTP {status} {resp_data}"
            print(f"  [X] FAILURE: {err_msg}")
            verification_failures.append(err_msg)
            continue

        # Verify force_password_change is returned
        fpc = resp_data.get("force_password_change")
        if fpc is not True:
            err_msg = f"Account {role} ID {user_id} login did not return force_password_change=True: {resp_data}"
            print(f"  [X] FAILURE: {err_msg}")
            verification_failures.append(err_msg)
            continue

        item["verified_login"] = True
        verified_count += 1
        if idx % 20 == 0 or idx == 128:
            print(f"  Verified {idx}/128 accounts successfully...")

    if verification_failures:
        print(f"\n[CRITICAL ERROR] {len(verification_failures)} accounts failed verification!")
        for fail in verification_failures:
            print(f"  - {fail}")
        raise RuntimeError("Credential verification failed! Stop and investigate.")

    print(f"\n  [SUCCESS] All {verified_count}/128 accounts verified with HTTP 200 and force_password_change=True!")

    # STEP 5: VERIFY DATABASE STATE
    print("\n[STEP 5] Verifying final database state...")
    db.expire_all()
    recheck_users = db.query(User).filter(User.is_active == True).all()

    assert len(recheck_users) == 128
    assert all(u.force_password_change is True for u in recheck_users), "Some accounts do not have force_password_change=True!"
    assert all(u.failed_login_attempts == 0 for u in recheck_users), "Some accounts have non-zero failed_login_attempts!"
    assert all(u.locked_until is None for u in recheck_users), "Some accounts are locked!"
    print("  Database verified:")
    print("  - 128 accounts have force_password_change = True")
    print("  - 128 accounts have failed_login_attempts = 0")
    print("  - 128 accounts have locked_until = None")

    # STEP 6: CREATE CREDENTIAL INVENTORY EXCEL FILE
    print(f"\n[STEP 6] Generating Excel inventory file: {EXCEL_OUTPUT_PATH}...")
    wb = openpyxl.Workbook()

    # Style presets
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    mono_font = Font(name="Consolas", size=10, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # SHEET 1: ALL ACCOUNTS
    ws1 = wb.active
    ws1.title = "ALL ACCOUNTS"
    headers1 = [
        "User ID", "Role", "Full Name", "Username", "Email",
        "Student Roll Number", "Faculty Employee ID", "New Password",
        "Force Password Change", "Verified Login", "Generated At"
    ]
    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for row_idx, item in enumerate(account_records, start=2):
        row_vals = [
            item["user_id"],
            item["role"],
            item["full_name"],
            item["username"],
            item["email"],
            item["roll_number"],
            item["employee_id"],
            item["new_password"],
            "TRUE" if item["force_password_change"] else "FALSE",
            "YES" if item["verified_login"] else "NO",
            item["generated_at"]
        ]
        ws1.append(row_vals)
        for col_idx in range(1, len(row_vals) + 1):
            c = ws1.cell(row=row_idx, column=col_idx)
            c.border = border_thin
            if col_idx in [1, 2, 6, 7, 9, 10, 11]:
                c.alignment = center_align
                c.font = data_font
            elif col_idx == 8:  # Password column
                c.alignment = center_align
                c.font = mono_font
            else:
                c.alignment = left_align
                c.font = data_font

    # Adjust column widths for Sheet 1
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # SHEET 2: SUMMARY
    ws2 = wb.create_sheet(title="SUMMARY")
    ws2.append(["Metric", "Count / Value", "Notes"])
    for col_idx in range(1, 4):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    summary_rows = [
        ("Total Accounts", len(account_records), "All active accounts in attendance_system.db"),
        ("Students", len(student_users), "Class 1 (64) + Class 2 (55) students"),
        ("Faculty", len(faculty_users), "Official faculty accounts (FAC01-FAC08)"),
        ("Admin", len(admin_users), "System administrator account"),
        ("Successfully Rotated", len(account_records), "Hashed with Bcrypt and committed to database"),
        ("Successfully Verified", verified_count, "Verified via live authentication endpoints with HTTP 200"),
        ("Failed Rotations", 0, "No failures encountered"),
        ("Locked Accounts", 0, "All accounts active with locked_until = None"),
        ("Accounts Requiring Manual Review", 0, "None"),
        ("Password Complexity", "14 chars, mixed-case, digits, symbols", "Guarantees !@#$%^&* and high entropy"),
        ("Rotated Timestamp", generation_time_str, "Coordinated credential rotation")
    ]

    for row_idx, r in enumerate(summary_rows, start=2):
        ws2.append(list(r))
        for col_idx in range(1, 4):
            c = ws2.cell(row=row_idx, column=col_idx)
            c.border = border_thin
            c.font = data_font
            if col_idx == 2:
                c.alignment = center_align
            else:
                c.alignment = left_align

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 5, 16)

    # SHEET 3: INSTRUCTIONS
    ws3 = wb.create_sheet(title="INSTRUCTIONS")
    ws3.append(["Item #", "Security Guideline / Instructions"])
    for col_idx in range(1, 3):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    instructions_lines = [
        (1, "These credentials are TEMPORARY passwords intended for the initial system rollout or security rotation."),
        (2, "All accounts have force_password_change = TRUE configured. Users will be automatically prompted to set a secure permanent password upon first login."),
        (3, "The permanent password chosen by the user must be at least 8 characters long, distinct from the temporary password, and confirmed."),
        (4, "DO NOT distribute this spreadsheet publicly or commit it to version control (Git)."),
        (5, "Distribute credentials to individual students and faculty via authorized private college administrative channels only."),
        (6, "Once credentials have been safely handed over or distributed, store this master spreadsheet in an encrypted location or securely delete it.")
    ]

    for row_idx, line in enumerate(instructions_lines, start=2):
        ws3.append(list(line))
        for col_idx in range(1, 3):
            c = ws3.cell(row=row_idx, column=col_idx)
            c.border = border_thin
            c.font = data_font
            if col_idx == 1:
                c.alignment = center_align
            else:
                c.alignment = left_align

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(EXCEL_OUTPUT_PATH)
    print(f"  Excel file written successfully to: {EXCEL_OUTPUT_PATH}")

    db.close()
    print("\n" + "=" * 75)
    print("ROTATION & VERIFICATION COMPLETED WITH 100% SUCCESS")
    print("=" * 75)

if __name__ == "__main__":
    execute_rotation()
