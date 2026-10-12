import os
import sys
import time
import subprocess
import openpyxl

ADB = r"C:\Users\Abishek\AppData\Local\Android\Sdk\platform-tools\adb.exe"
EXCEL_PATH = r"C:\Users\Abishek\Documents\Attendance_Marker\ALL_ACCOUNT_CREDENTIALS_20260926.xlsx"
OUTPUT_DIR = r"C:\Users\Abishek\.gemini\antigravity\brain\7f90dbc7-1f76-43d0-a6a4-acccb732defd"

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb["Faculty"]
pw = None
for row in range(5, ws.max_row + 1):
    u = str(ws.cell(row, 4).value or "").strip().lower()
    if u == "fac05":
        pw = str(ws.cell(row, 5).value or "").strip()
        break

print("Found pw for fac05:", bool(pw))

def adb_cmd(cmd):
    return subprocess.run(f'"{ADB}" -s FY6PHIEA69AMQ8HU {cmd}', shell=True, capture_output=True, text=True)

# Tap username field
print("Tapping username...")
adb_cmd("shell input tap 540 1292")
time.sleep(0.5)
# Select all and delete
adb_cmd("shell input keyevent 29 --meta 114") # Ctrl+A
adb_cmd("shell input keyevent 67") # Backspace
for _ in range(15):
    adb_cmd("shell input keyevent 67")
adb_cmd("shell input text fac05")
time.sleep(0.5)

# Press TAB to focus password field
print("TAB to password...")
adb_cmd("shell input keyevent 61")
time.sleep(0.5)
for _ in range(30):
    adb_cmd("shell input keyevent 67")

# Use direct argument call for password
escaped = pw.replace("&", "\\&").replace(" ", "%s").replace("'", "\\'").replace('"', '\\"')
subprocess.run([ADB, "-s", "FY6PHIEA69AMQ8HU", "shell", "input", "text", escaped], check=True)
time.sleep(0.5)

# Dismiss keyboard (back key)
adb_cmd("shell input keyevent 4")
time.sleep(0.8)

# Tap Faculty Sign In button at [60,1688][1020,1832] -> (540, 1760)
print("Tapping Faculty Sign In...")
adb_cmd("shell input tap 540 1760")
time.sleep(4.0)

# Capture screen
adb_cmd("shell screencap -p /sdcard/sc.png")
adb_cmd(f'pull /sdcard/sc.png "{OUTPUT_DIR}\\15_faculty_dashboard_fixed.png"')
print("Captured 15_faculty_dashboard_fixed.png")
