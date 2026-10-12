import os
import sys
import time
import subprocess
import openpyxl

ADB = r"C:\Users\Abishek\AppData\Local\Android\Sdk\platform-tools\adb.exe"
EXCEL_PATH = r"C:\Users\Abishek\Documents\Attendance_Marker\ALL_ACCOUNT_CREDENTIALS_20260926.xlsx"
OUTPUT_DIR = r"C:\Users\Abishek\.gemini\antigravity\brain\7f90dbc7-1f76-43d0-a6a4-acccb732defd"

def adb(cmd):
    full_cmd = f'"{ADB}" {cmd}'
    res = subprocess.run(full_cmd, shell=True, capture_output=True, text=True)
    return res.stdout.strip()

def tap(x, y, delay=1.0):
    adb(f"shell input tap {x} {y}")
    time.sleep(delay)

def keyevent(code, delay=0.5):
    adb(f"shell input keyevent {code}")
    time.sleep(delay)

def type_text(text, delay=0.5):
    escaped = text.replace("&", "\\&").replace(" ", "%s").replace("'", "\\'").replace('"', '\\"')
    adb(f"shell input text '{escaped}'")
    time.sleep(delay)

def swipe(x1, y1, x2, y2, duration=300, delay=1.0):
    adb(f"shell input swipe {x1} {y1} {x2} {y2} {duration}")
    time.sleep(delay)

def capture_screen(filename):
    device_path = f"/sdcard/{filename}"
    local_path = os.path.join(OUTPUT_DIR, filename)
    adb(f"shell screencap -p {device_path}")
    adb(f'pull {device_path} "{local_path}"')
    print(f"Captured: {local_path}")
    return local_path

def get_student_password(reg_no):
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["Students"]
    for row in range(5, ws.max_row + 1):
        r_no = str(ws.cell(row, 2).value or "").strip()
        if r_no == reg_no:
            return str(ws.cell(row, 5).value or "").strip()
    return None

def get_faculty_password(username):
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["Faculty"]
    for row in range(5, ws.max_row + 1):
        u = str(ws.cell(row, 4).value or "").strip().lower()
        if u == username.lower():
            return str(ws.cell(row, 5).value or "").strip()
    return None

def run_class1_flow():
    print("=== STARTING CLASS 1 STUDENT FLOW ===")
    pw = get_student_password("721025149001")
    if not pw:
        print("ERROR: Password not found for 721025149001")
        return False
    
    # Ensure app is on login screen
    adb("shell am force-stop com.aifaceattendance.mobile_app")
    time.sleep(1.0)
    adb("shell am start -n com.aifaceattendance.mobile_app/.MainActivity")
    time.sleep(3.5)
    
    capture_screen("00_login_screen.png")

    # 1. Tap Register Number field (center 540, 1082)
    print("Tapping Register Number field...")
    tap(540, 1082, delay=0.8)
    type_text("721025149001", delay=0.5)
    keyevent(4, delay=0.8) # Dismiss keyboard

    # 2. Tap Password field (center 540, 1316)
    print("Tapping Password field...")
    tap(540, 1316, delay=0.8)
    type_text(pw, delay=0.5)
    keyevent(4, delay=0.8) # Dismiss keyboard

    # 3. Tap Student Sign In (center 540, 1520)
    print("Tapping Student Sign In...")
    tap(540, 1520, delay=4.0)

    # 4. Capture Dashboard
    print("Capturing 01_class1_dashboard.png...")
    capture_screen("01_class1_dashboard.png")

    # 5. Tap Timetable bottom nav tab (x = 540, y = 2300)
    print("Tapping Timetable tab...")
    tap(540, 2300, delay=2.5)
    capture_screen("02_class1_timetable_today.png")

    # 6. Tap "THU" in day selection bar (x = 630, y = 270)
    print("Tapping THU day button...")
    tap(630, 270, delay=2.0)
    capture_screen("03_class1_timetable_thu_iks.png")

    # 7. Tap "SAT" in day selection bar (x = 990, y = 270)
    print("Tapping SAT day button...")
    tap(990, 270, delay=2.0)
    capture_screen("04_class1_timetable_sat_iks.png")

    # 8. Tap Profile tab (x = 900, y = 2300)
    print("Tapping Profile tab...")
    tap(900, 2300, delay=2.5)
    capture_screen("05_class1_profile.png")

    # Scroll down on Profile to show Actions
    print("Scrolling down on Profile...")
    swipe(540, 1800, 540, 800, duration=300, delay=1.5)
    capture_screen("06_class1_profile_actions.png")

    # Tap Sign Out
    print("Signing out...")
    tap(540, 1998, delay=1.0)
    tap(540, 1483, delay=2.0)
    return True

def run_class2_flow():
    print("=== STARTING CLASS 2 STUDENT FLOW ===")
    pw = get_student_password("721024149001")
    if not pw:
        print("ERROR: Password not found for 721024149001")
        return False
    
    # 1. Tap Register Number field (center 540, 1082)
    print("Tapping Register Number field...")
    tap(540, 1082, delay=0.8)
    type_text("721024149001", delay=0.5)
    keyevent(4, delay=0.8) # Dismiss keyboard

    # 2. Tap Password field (center 540, 1316)
    print("Tapping Password field...")
    tap(540, 1316, delay=0.8)
    type_text(pw, delay=0.5)
    keyevent(4, delay=0.8) # Dismiss keyboard

    # 3. Tap Student Sign In (center 540, 1520)
    print("Tapping Student Sign In...")
    tap(540, 1520, delay=4.0)

    # 4. Capture Dashboard
    print("Capturing 09_class2_dashboard.png...")
    capture_screen("09_class2_dashboard.png")

    # 5. Tap Timetable bottom nav tab (x = 540, y = 2316)
    print("Tapping Timetable tab...")
    tap(540, 2316, delay=2.5)
    capture_screen("10_class2_timetable_mon.png")

    # 6. Tap "THU" in day selection bar (x = 651, y = 381)
    print("Tapping THU day button...")
    tap(651, 381, delay=2.0)
    capture_screen("11_class2_timetable_thu.png")

    # 7. Tap Profile tab (x = 900, y = 2316)
    print("Tapping Profile tab...")
    tap(900, 2316, delay=2.5)
    capture_screen("12_class2_profile.png")

    # Scroll down on Profile to show Actions
    print("Scrolling down on Profile...")
    swipe(540, 1800, 540, 800, duration=300, delay=1.5)
    capture_screen("13_class2_profile_actions.png")

    # 8. Tap Sign Out and confirm
    print("Signing out...")
    tap(540, 1998, delay=1.0)
    tap(540, 1483, delay=2.0)
    capture_screen("14_login_after_class2.png")

    return True

def run_faculty_flow():
    print("=== STARTING FACULTY FLOW (fac05 / Santhiya) ===")
    pw = get_faculty_password("fac05")
    if not pw:
        print("ERROR: Password not found for fac05")
        return False

    # Close error banner if open
    print("Dismissing any error banner...")
    tap(915, 1067, delay=0.5)

    # 1. Tap Faculty tab on login screen (x = 778, y = 866)
    print("Tapping Faculty tab...")
    tap(778, 866, delay=1.0)

    # 2. Tap Username field (center 540, 1292)
    print("Tapping Username field...")
    tap(540, 1292, delay=0.8)
    type_text("fac05", delay=0.5)
    keyevent(4, delay=0.8) # Dismiss keyboard

    # 3. Tap Password field (center 540, 1496)
    print("Tapping Password field...")
    tap(540, 1496, delay=0.8)
    type_text(pw, delay=0.5)
    keyevent(4, delay=0.8) # Dismiss keyboard

    # 4. Tap Faculty Sign In (center 540, 1700)
    print("Tapping Faculty Sign In...")
    tap(540, 1700, delay=4.5)

    # 5. Capture Faculty Dashboard
    print("Capturing 15_faculty_dashboard.png...")
    capture_screen("15_faculty_dashboard.png")

    # 6. Dump hierarchy of Faculty screen to get schedule day selector & navigation
    adb("shell uiautomator dump /sdcard/fac_screen_dump.xml")
    adb(r'pull /sdcard/fac_screen_dump.xml "scripts/fac_screen_dump.xml"')

    return True

if __name__ == "__main__":
    action = sys.argv[1] if len(sys.argv) > 1 else "class1"
    if action == "class1":
        run_class1_flow()
    elif action == "class2":
        run_class2_flow()
    elif action == "faculty":
        run_faculty_flow()
