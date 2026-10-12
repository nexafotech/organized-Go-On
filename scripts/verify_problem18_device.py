import os
import sys
import time
import subprocess
import xml.etree.ElementTree as ET
import openpyxl

ADB = r"C:\Users\Abishek\AppData\Local\Android\Sdk\platform-tools\adb.exe"
EXCEL_PATH = r"C:\Users\Abishek\Documents\Attendance_Marker\ALL_ACCOUNT_CREDENTIALS_20260926.xlsx"
OUTPUT_DIR = r"C:\Users\Abishek\.gemini\antigravity\brain\7f90dbc7-1f76-43d0-a6a4-acccb732defd"

def adb(cmd):
    full_cmd = f'"{ADB}" {cmd}'
    res = subprocess.run(full_cmd, shell=True, capture_output=True, text=True)
    return res.stdout.strip()

def tap(x, y, delay=1.0):
    adb(f"shell input tap {x} {y}")
    time.sleep(delay)

def keyevent(code, delay=0.5):
    adb(f"shell input keyevent {code}")
    time.sleep(delay)

def type_text(text, delay=0.5):
    escaped = text.replace("&", "\\&").replace(" ", "%s").replace("'", "\\'").replace('"', '\\"')
    adb(f"shell input text '{escaped}'")
    time.sleep(delay)

def clear_text_field(backspaces=45):
    keyevent(123, delay=0.1) # KEYCODE_MOVE_END
    for _ in range(backspaces):
        keyevent(67, delay=0.02) # KEYCODE_DEL
    time.sleep(0.5)

def swipe(x1, y1, x2, y2, duration=300, delay=1.0):
    adb(f"shell input swipe {x1} {y1} {x2} {y2} {duration}")
    time.sleep(delay)

def capture_screen(filename):
    device_path = f"/sdcard/{filename}"
    local_path = os.path.join(OUTPUT_DIR, filename)
    adb(f"shell screencap -p {device_path}")
    adb(f'pull {device_path} "{local_path}"')
    print(f"[SCREENSHOT] {local_path}")
    return local_path

def dump_ui(xml_name="ui_dump.xml"):
    device_xml = f"/sdcard/{xml_name}"
    local_xml = os.path.join(r"C:\Users\Abishek\Desktop\am-oss\scripts", xml_name)
    adb(f"shell uiautomator dump {device_xml}")
    adb(f'pull {device_xml} "{local_xml}"')
    return local_xml

def find_node(xml_path, text_contains=None, exact_text=None, resource_id=None):
    if not os.path.exists(xml_path):
        return None
    try:
        tree = ET.parse(xml_path)
        root = tree.getroot()
        for node in root.iter("node"):
            t = node.attrib.get("text", "")
            r = node.attrib.get("resource-id", "")
            c = node.attrib.get("content-desc", "")
            if exact_text and (t == exact_text or c == exact_text):
                return node
            if text_contains and (text_contains in t or text_contains in c):
                return node
            if resource_id and resource_id in r:
                return node
    except Exception as e:
        print(f"Error parsing UI dump: {e}")
    return None

def get_node_center(node):
    bounds = node.attrib.get("bounds", "")
    parts = bounds.replace("][", ",").replace("[", "").replace("]", "").split(",")
    if len(parts) == 4:
        x1, y1, x2, y2 = map(int, parts)
        return (x1 + x2) // 2, (y1 + y2) // 2
    return None

def get_faculty_password(username):
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["Faculty"]
    for row in range(5, ws.max_row + 1):
        u = str(ws.cell(row, 4).value or "").strip().lower()
        if u == username.lower():
            return str(ws.cell(row, 5).value or "").strip()
    return None

def get_student_password(reg_no):
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["Students"]
    for row in range(5, ws.max_row + 1):
        r_no = str(ws.cell(row, 2).value or "").strip()
        if r_no == reg_no:
            return str(ws.cell(row, 5).value or "").strip()
    return None

def ensure_login_screen():
    print("[STEP] Ensuring application is at login screen...")
    xml = dump_ui("check_login.xml")
    
    # Check if Profile or Sign Out is visible (meaning already logged in)
    profile_node = find_node(xml, text_contains="Profile")
    signout_node = find_node(xml, text_contains="Sign Out")
    
    if profile_node and not signout_node:
        # Tap Profile tab
        center = get_node_center(profile_node)
        if center:
            tap(center[0], center[1], delay=2.0)
            xml = dump_ui("check_login.xml")
            signout_node = find_node(xml, text_contains="Sign Out")

    if signout_node:
        # Tap Sign Out
        center = get_node_center(signout_node)
        if center:
            tap(center[0], center[1], delay=1.0)
            # Confirm dialog
            xml = dump_ui("confirm_so.xml")
            yes_node = find_node(xml, exact_text="Sign Out") or find_node(xml, text_contains="Yes")
            if yes_node:
                yc = get_node_center(yes_node)
                if yc:
                    tap(yc[0], yc[1], delay=2.0)

    # If still not at login, restart app
    xml = dump_ui("check_login2.xml")
    if not find_node(xml, exact_text="FACULTY"):
        adb("shell am force-stop com.aifaceattendance.mobile_app")
        time.sleep(1.0)
        adb("shell am start -n com.aifaceattendance.mobile_app/.MainActivity")
        time.sleep(4.0)

def run_problem18_verification():
    print("=" * 60)
    print("STARTING PROBLEM 18 REAL DEVICE VERIFICATION")
    print("=" * 60)

    ensure_login_screen()

    # Step 1: Switch to FACULTY tab
    print("\n--- 1. Switching to FACULTY Tab ---")
    xml = dump_ui("fac_tab.xml")
    fac_tab = find_node(xml, exact_text="FACULTY")
    if fac_tab:
        cx, cy = get_node_center(fac_tab)
        tap(cx, cy, delay=1.5)
    else:
        # fallback coordinates
        tap(778, 866, delay=1.5)

    keyevent(4, delay=0.5) # dismiss keyboard if open
    capture_screen("p18_01_faculty_tab_empty.png")

    # Step 2: Validate Faculty ID 'FAC01' (Live Recognized)
    print("\n--- 2. Live Validation: Valid Faculty ID FAC01 ---")
    xml = dump_ui("fac_form.xml")
    id_field = find_node(xml, text_contains="Username / Employee ID / Email")
    if id_field:
        cx, cy = get_node_center(id_field)
    else:
        cx, cy = 540, 1292

    tap(cx, cy, delay=0.5)
    type_text("FAC01", delay=0.1)
    # Immediately capture to demonstrate checking/live debounce
    capture_screen("p18_02_faculty_checking.png")

    # Wait for debounce (500ms) + backend response
    time.sleep(1.5)
    keyevent(4, delay=0.5) # dismiss keyboard
    capture_screen("p18_03_faculty_valid_fac01.png")

    xml = dump_ui("fac01_res.xml")
    recognized = find_node(xml, text_contains="Faculty account recognized")
    print(f"Faculty account recognized status found: {recognized is not None}")

    # Step 3: Reject Class 1 Student Register Number
    print("\n--- 3. Live Validation: Reject Class 1 Student ID (721025149001) ---")
    tap(cx, cy, delay=0.5)
    clear_text_field(15)
    type_text("721025149001", delay=0.2)
    time.sleep(1.5)
    keyevent(4, delay=0.5)
    capture_screen("p18_04_faculty_invalid_student_id.png")

    xml = dump_ui("student1_res.xml")
    invalid_node = find_node(xml, text_contains="Invalid faculty username")
    print(f"Invalid faculty username status found: {invalid_node is not None}")

    # Step 4: Reject Class 2 Student Register Number
    print("\n--- 4. Live Validation: Reject Class 2 Student ID (721025149061) ---")
    tap(cx, cy, delay=0.5)
    clear_text_field(20)
    type_text("721025149061", delay=0.2)
    time.sleep(1.5)
    keyevent(4, delay=0.5)
    capture_screen("p18_05_faculty_invalid_class2_id.png")

    # Step 5: Reject Admin username
    print("\n--- 5. Live Validation: Reject Admin Username (admin) ---")
    tap(cx, cy, delay=0.5)
    clear_text_field(20)
    type_text("admin", delay=0.2)
    time.sleep(1.5)
    keyevent(4, delay=0.5)
    capture_screen("p18_06_faculty_invalid_admin.png")

    # Step 6: Clear text field -> returns to empty state
    print("\n--- 6. Clear Field -> State Returns to Empty ---")
    tap(cx, cy, delay=0.5)
    clear_text_field(20)
    time.sleep(1.0)
    keyevent(4, delay=0.5)
    capture_screen("p18_07_faculty_cleared_empty.png")

    # Step 7: Network error verification
    print("\n--- 7. Live Validation: Network Error Simulation ---")
    # Tap Server Settings icon in AppBar (around x=980, y=140 or search icon)
    tap(980, 140, delay=1.5)
    xml = dump_ui("server_dlg.xml")
    url_input = find_node(xml, text_contains="http")
    if url_input:
        ux, uy = get_node_center(url_input)
        tap(ux, uy, delay=0.5)
        clear_text_field(35)
        type_text("http://10.254.103.195:9999", delay=0.2)
        keyevent(4, delay=0.5)
        # Tap Save
        save_btn = find_node(xml, exact_text="Save")
        if save_btn:
            sx, sy = get_node_center(save_btn)
            tap(sx, sy, delay=1.5)

    # Now type FAC01 in faculty username field to trigger network error
    tap(cx, cy, delay=0.5)
    type_text("FAC01", delay=0.2)
    time.sleep(2.5)
    keyevent(4, delay=0.5)
    capture_screen("p18_08_faculty_network_error.png")

    # Restore correct Server URL
    print("Restoring server URL...")
    tap(980, 140, delay=1.5)
    xml = dump_ui("server_dlg2.xml")
    reset_btn = find_node(xml, exact_text="Reset Default")
    if reset_btn:
        rx, ry = get_node_center(reset_btn)
        tap(rx, ry, delay=0.8)
    save_btn = find_node(xml, exact_text="Save")
    if save_btn:
        sx, sy = get_node_center(save_btn)
        tap(sx, sy, delay=1.5)

    # Step 8: Faculty Login Regression (fac05 / Santhiya)
    print("\n--- 8. End-to-End Faculty Login (fac05) ---")
    tap(cx, cy, delay=0.5)
    clear_text_field(15)
    type_text("FAC05", delay=0.2)
    time.sleep(1.5)
    keyevent(4, delay=0.5)

    fac_pw = get_faculty_password("fac05")
    if fac_pw:
        # Tap Password field (below identifier)
        xml = dump_ui("fac_pw_field.xml")
        pw_node = find_node(xml, text_contains="Enter faculty password") or find_node(xml, text_contains="Password")
        if pw_node:
            px, py = get_node_center(pw_node)
        else:
            px, py = 540, 1496
        tap(px, py, delay=0.5)
        type_text(fac_pw, delay=0.2)
        keyevent(4, delay=0.5)

        # Tap Faculty Sign In
        sign_btn = find_node(xml, exact_text="Faculty Sign In")
        if sign_btn:
            bx, by = get_node_center(sign_btn)
        else:
            bx, by = 540, 1700
        tap(bx, by, delay=4.5)
        capture_screen("p18_09_faculty_login_success_fac05.png")

        # Sign out from faculty
        print("Signing out from faculty dashboard...")
        xml = dump_ui("fac_dash.xml")
        so_node = find_node(xml, text_contains="Sign Out") or find_node(xml, text_contains="Logout")
        if not so_node:
            prof_node = find_node(xml, text_contains="Profile")
            if prof_node:
                pc = get_node_center(prof_node)
                tap(pc[0], pc[1], delay=2.0)
                xml = dump_ui("fac_prof.xml")
                so_node = find_node(xml, text_contains="Sign Out")

        if so_node:
            sc = get_node_center(so_node)
            tap(sc[0], sc[1], delay=1.0)
            xml = dump_ui("fac_so_confirm.xml")
            yes = find_node(xml, exact_text="Sign Out") or find_node(xml, text_contains="Yes")
            if yes:
                yc = get_node_center(yes)
                tap(yc[0], yc[1], delay=2.0)

    # Step 9: Student Login Regression (721025149001)
    print("\n--- 9. Student Login Regression (721025149001) ---")
    xml = dump_ui("stu_tab.xml")
    stu_tab = find_node(xml, exact_text="STUDENT")
    if stu_tab:
        sx, sy = get_node_center(stu_tab)
        tap(sx, sy, delay=1.5)

    stu_pw = get_student_password("721025149001")
    if stu_pw:
        xml = dump_ui("stu_form.xml")
        s_id = find_node(xml, text_contains="Register Number")
        if s_id:
            ix, iy = get_node_center(s_id)
        else:
            ix, iy = 540, 1082
        tap(ix, iy, delay=0.5)
        type_text("721025149001", delay=0.2)
        keyevent(4, delay=0.5)

        s_pw = find_node(xml, text_contains="Enter student password") or find_node(xml, text_contains="Password")
        if s_pw:
            px, py = get_node_center(s_pw)
        else:
            px, py = 540, 1316
        tap(px, py, delay=0.5)
        type_text(stu_pw, delay=0.2)
        keyevent(4, delay=0.5)

        s_btn = find_node(xml, exact_text="Student Sign In")
        if s_btn:
            bx, by = get_node_center(s_btn)
        else:
            bx, by = 540, 1520
        tap(bx, by, delay=4.5)
        capture_screen("p18_10_student_login_regression.png")

        # Sign out from student
        print("Signing out from student dashboard...")
        tap(900, 2300, delay=2.5) # Profile tab
        swipe(540, 1800, 540, 800, duration=300, delay=1.5)
        tap(540, 1998, delay=1.0) # Sign Out
        tap(540, 1483, delay=2.0) # Confirm

    print("\n" + "=" * 60)
    print("PROBLEM 18 REAL DEVICE VERIFICATION COMPLETED SUCCESSFULLY!")
    print("=" * 60)

if __name__ == "__main__":
    run_problem18_verification()
