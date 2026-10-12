"""
Milestone 15F - System-Wide Secure Credential Regeneration and Full Backend Verification
Authoritative Execution Script

CRITICAL CONSTRAINTS:
- Verified 100% against live FastAPI backend.
- Mobile Flutter UI verification skipped (unavailable/runnable).
- 0 plaintext passwords logged to stdout, logs, or stored in database.
- Excel file saved strictly outside Git repository.
- Admin account credentials preserved intact.
"""

import os
import sys
import shutil
import sqlite3
import datetime
from datetime import timezone
import json
import urllib.request
import urllib.error

# Add am-oss root to path
REPO_PATH = r"C:\Users\Abishek\Desktop\am-oss"
if REPO_PATH not in sys.path:
    sys.path.insert(0, REPO_PATH)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.security import get_password_hash, verify_password
from app.services.class1_onboarding_service import generate_secure_temporary_password

DB_PATH = os.path.join(REPO_PATH, "attendance_system.db")
OUTPUT_DIR = r"C:\Users\Abishek\Documents\Attendance_Marker"
CANONICAL_EXCEL_FILENAME = "ALL_ACCOUNT_CREDENTIALS_20260926.xlsx"
BACKEND_BASE_URL = "http://127.0.0.1:8000"


def http_post(url: str, payload: dict, token: str = None):
    data_bytes = json.dumps(payload).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, data=data_bytes, headers=headers, method="POST")
    try:
        with urllib.request.urlopen(req) as resp:
            body = resp.read().decode("utf-8")
            return resp.status, json.loads(body) if body else {}
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8")
        try:
            parsed = json.loads(body)
        except Exception:
            parsed = {"detail": body}
        return e.code, parsed


def http_get(url: str, token: str = None):
    headers = {}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, headers=headers, method="GET")
    try:
        with urllib.request.urlopen(req) as resp:
            body = resp.read().decode("utf-8")
            return resp.status, json.loads(body) if body else {}
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8")
        try:
            parsed = json.loads(body)
        except Exception:
            parsed = {"detail": body}
        return e.code, parsed


def main():
    print("=" * 70)
    print("SECURE CREDENTIAL REGENERATION & 100% BACKEND VERIFICATION")
    print("=" * 70)

    # -------------------------------------------------------------
    # 1. PRE-FLIGHT DATABASE & BACKEND HEALTH AUDIT
    # -------------------------------------------------------------
    print("\n[Step 1] Pre-Flight Database & Live Backend Audit...")
    
    # Check backend health
    try:
        status, health_data = http_get(f"{BACKEND_BASE_URL}/api/v1/health")
        if status != 200:
            raise RuntimeError(f"Backend health check failed: HTTP {status} {health_data}")
        print(f"  Live Backend Status: {health_data.get('status')} (HTTP {status})")
    except Exception as e:
        raise RuntimeError(f"Cannot connect to live backend at {BACKEND_BASE_URL}: {e}")

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    cur.execute("SELECT count(*) FROM students")
    pre_students = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM faculty")
    pre_faculty = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM users")
    pre_users = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM class_sections")
    pre_sections = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM subjects")
    pre_subjects = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM attendance_records")
    pre_attendance = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM class_sessions")
    pre_sessions = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM face_embeddings")
    pre_embeddings = cur.fetchone()[0]

    # Class breakdown
    cur.execute("SELECT s.class_id, cs.name, count(*) FROM students s JOIN class_sections cs ON s.class_id = cs.id GROUP BY s.class_id ORDER BY s.class_id")
    class_breakdown = cur.fetchall()

    print(f"  Students Total: {pre_students} (Expected: 119)")
    for cid, cname, cnt in class_breakdown:
        print(f"    - Section {cid} ({cname}): {cnt} students")
    print(f"  Faculty Total: {pre_faculty} (Expected: 8)")
    print(f"  Users Total: {pre_users} (Expected: 128)")
    print(f"  Class Sections: {pre_sections} (Expected: 2)")
    print(f"  Subjects: {pre_subjects} (Expected: 27)")
    print(f"  Attendance Records: {pre_attendance} (Expected: 0)")
    print(f"  Class Sessions: {pre_sessions} (Expected: 0)")
    print(f"  Face Embeddings: {pre_embeddings} (Expected: 0)")

    if pre_students != 119 or pre_faculty != 8 or pre_users != 128:
        raise ValueError(f"Pre-flight audit failed: Invariant count mismatch! (students={pre_students}, faculty={pre_faculty}, users={pre_users})")
    if pre_attendance != 0 or pre_sessions != 0 or pre_embeddings != 0:
        raise ValueError("Pre-flight audit failed: Non-zero attendance/session/biometric records found!")

    # Verify admin account
    cur.execute("SELECT id, username, email, role, is_active FROM users WHERE role = 'ADMIN'")
    admin_row = cur.fetchone()
    if not admin_row or admin_row[1] != "admin":
        raise ValueError(f"Admin account unexpected: {admin_row}")
    admin_id, admin_user, admin_email, admin_role, admin_active = admin_row
    print(f"  Admin User Verified: ID={admin_id}, Username='{admin_user}' (Unmodified)")

    # -------------------------------------------------------------
    # 2. CREATE DATABASE BACKUPS
    # -------------------------------------------------------------
    print("\n[Step 2] Creating Database Backups...")
    ts_now = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_db_local = os.path.join(REPO_PATH, f"attendance_system_backup_credential_regen_{ts_now}.db")
    shutil.copy2(DB_PATH, backup_db_local)
    print(f"  Local DB Backup created: {backup_db_local}")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    backup_db_doc = os.path.join(OUTPUT_DIR, f"attendance_system_backup_credential_regen_{ts_now}.db")
    shutil.copy2(DB_PATH, backup_db_doc)
    print(f"  Documents DB Backup created: {backup_db_doc}")

    # -------------------------------------------------------------
    # 3. RETRIEVE ALL ACCOUNTS & GENERATE SECURE PASSWORDS
    # -------------------------------------------------------------
    print("\n[Step 3] Generating Cryptographically Secure Passwords...")
    
    # Students
    cur.execute("""
        SELECT s.id, s.user_id, s.roll_number, s.full_name, cs.name as class_name, u.username, u.is_active, s.class_id
        FROM students s
        JOIN users u ON s.user_id = u.id
        JOIN class_sections cs ON s.class_id = cs.id
        ORDER BY s.class_id ASC, s.roll_number ASC
    """)
    student_records = cur.fetchall()
    print(f"  Retrieved {len(student_records)} student accounts.")

    # Faculty
    cur.execute("""
        SELECT f.id, f.user_id, f.employee_id, f.full_name, f.department, u.username, u.is_active
        FROM faculty f
        JOIN users u ON f.user_id = u.id
        ORDER BY f.employee_id ASC
    """)
    faculty_records = cur.fetchall()
    print(f"  Retrieved {len(faculty_records)} faculty accounts.")

    generated_passwords = set()
    student_cred_list = []
    faculty_cred_list = []

    # Generate for Faculty (FAC01-FAC08)
    for f_id, u_id, emp_id, full_name, department, username, is_active in faculty_records:
        while True:
            pwd = generate_secure_temporary_password(8)
            if pwd not in generated_passwords:
                generated_passwords.add(pwd)
                break
        
        # Verify complexity
        assert len(pwd) == 8
        assert any(c.isupper() for c in pwd)
        assert any(c.islower() for c in pwd)
        assert any(c.isdigit() for c in pwd)
        assert any(c in "!@#$%^&*" for c in pwd)

        hashed = get_password_hash(pwd)
        item = {
            "type": "FACULTY",
            "db_id": f_id,
            "user_id": u_id,
            "role": "FACULTY",
            "class_name": "N/A",
            "department": department or "CSE",
            "identifier": emp_id,
            "username": username,
            "full_name": full_name,
            "temp_password": pwd,
            "hashed_password": hashed,
            "force_password_change": "TRUE",
            "account_status": "ACTIVE" if is_active else "INACTIVE"
        }
        faculty_cred_list.append(item)

    # Generate for Students (119 students)
    for s_id, u_id, roll_no, full_name, class_name, username, is_active, class_id in student_records:
        while True:
            pwd = generate_secure_temporary_password(8)
            if pwd not in generated_passwords:
                generated_passwords.add(pwd)
                break
        
        # Verify complexity
        assert len(pwd) == 8
        assert any(c.isupper() for c in pwd)
        assert any(c.islower() for c in pwd)
        assert any(c.isdigit() for c in pwd)
        assert any(c in "!@#$%^&*" for c in pwd)

        hashed = get_password_hash(pwd)
        item = {
            "type": "STUDENT",
            "db_id": s_id,
            "user_id": u_id,
            "role": "STUDENT",
            "class_name": class_name,
            "department": "CSE",
            "identifier": roll_no,
            "username": username,
            "full_name": full_name,
            "temp_password": pwd,
            "hashed_password": hashed,
            "force_password_change": "TRUE",
            "account_status": "ACTIVE" if is_active else "INACTIVE",
            "class_id": class_id
        }
        student_cred_list.append(item)

    total_generated = len(generated_passwords)
    print(f"  Generated {total_generated} distinct passwords across 8 faculty and 119 students.")
    if total_generated != 127:
        raise ValueError(f"Expected 127 unique passwords, got {total_generated}")
    print("  Password Complexity & Uniqueness Check: PASS (127/127 unique, 8 chars, mixed-case, digits, symbols)")

    # -------------------------------------------------------------
    # 4. DATABASE UPDATE IN SAFE ATOMIC TRANSACTION
    # -------------------------------------------------------------
    print("\n[Step 4] Updating Database in Atomic Transaction...")
    cur.execute("BEGIN TRANSACTION")
    try:
        now_utc = datetime.datetime.now(timezone.utc).isoformat()

        # Update Faculty
        for item in faculty_cred_list:
            cur.execute("""
                UPDATE users
                SET hashed_password = ?,
                    force_password_change = 1,
                    session_token = NULL,
                    failed_login_attempts = 0,
                    locked_until = NULL
                WHERE id = ? AND role = 'FACULTY'
            """, (item["hashed_password"], item["user_id"]))
            if cur.rowcount != 1:
                raise RuntimeError(f"Failed to update user_id {item['user_id']} (role FACULTY)")

        # Update Students
        for item in student_cred_list:
            cur.execute("""
                UPDATE users
                SET hashed_password = ?,
                    force_password_change = 1,
                    session_token = NULL,
                    failed_login_attempts = 0,
                    locked_until = NULL
                WHERE id = ? AND role = 'STUDENT'
            """, (item["hashed_password"], item["user_id"]))
            if cur.rowcount != 1:
                raise RuntimeError(f"Failed to update user_id {item['user_id']} (role STUDENT)")

        # Insert audit log
        cur.execute("""
            INSERT INTO audit_logs (user_id, action, target, details, timestamp)
            VALUES (?, ?, ?, ?, ?)
        """, (
            admin_id,
            "BULK_TEMPORARY_PASSWORD_REGENERATION",
            "SYSTEM_USERS",
            "Bulk regeneration of secure temporary passwords for 119 students and 8 faculty accounts with forced password change.",
            now_utc
        ))

        # Check transaction state
        cur.execute("SELECT count(*) FROM users WHERE role = 'STUDENT' AND force_password_change = 1")
        fpc_students = cur.fetchone()[0]

        cur.execute("SELECT count(*) FROM users WHERE role = 'FACULTY' AND force_password_change = 1")
        fpc_faculty = cur.fetchone()[0]

        cur.execute("SELECT count(*) FROM users WHERE role = 'ADMIN' AND force_password_change = 1")
        fpc_admin = cur.fetchone()[0]

        if fpc_students != 119 or fpc_faculty != 8 or fpc_admin != 0:
            raise RuntimeError(f"Validation failed: fpc_students={fpc_students}, fpc_faculty={fpc_faculty}, fpc_admin={fpc_admin}")

        conn.commit()
        print("  Database Transaction Committed Successfully!")
        print(f"    - Student accounts updated: {fpc_students}")
        print(f"    - Faculty accounts updated: {fpc_faculty}")
        print(f"    - Admin account force_password_change: {fpc_admin} (Unmodified)")
    except Exception as e:
        conn.rollback()
        print(f"  DATABASE TRANSACTION ERROR: {e}")
        raise

    # -------------------------------------------------------------
    # 5. CREATE CONFIDENTIAL EXCEL WORKBOOK
    # -------------------------------------------------------------
    print("\n[Step 5] Creating Confidential Excel Spreadsheet...")
    
    canonical_excel_path = os.path.join(OUTPUT_DIR, CANONICAL_EXCEL_FILENAME)
    timestamped_excel_path = os.path.join(OUTPUT_DIR, f"ALL_ACCOUNT_CREDENTIALS_{ts_now}.xlsx")

    # Ensure output path is strictly OUTSIDE git repo
    repo_real = os.path.realpath(REPO_PATH)
    out_real = os.path.realpath(canonical_excel_path)
    if out_real.startswith(repo_real):
        raise ValueError(f"CRITICAL SAFETY VIOLATION: Excel target {out_real} is INSIDE Git repository {repo_real}!")
    print(f"  Path Safety Confirmed: {canonical_excel_path} is OUTSIDE git repository.")

    def build_workbook():
        wb = openpyxl.Workbook()
        default_sheet = wb.active

        # Styling
        font_title = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        fill_title = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")

        font_notice = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
        fill_notice = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

        font_data = Font(name="Calibri", size=11)
        font_password = Font(name="Consolas", size=11, bold=True, color="0F172A")
        fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # -----------------------------
        # Sheet 1: All Credentials
        # -----------------------------
        ws1 = wb.create_sheet(title="All Credentials")
        ws1.views.sheetView[0].showGridLines = True

        ws1.merge_cells("A1:I1")
        ws1["A1"].value = "CONFIDENTIAL — TEMPORARY LOGIN CREDENTIALS"
        ws1["A1"].font = font_title
        ws1["A1"].fill = fill_title
        ws1["A1"].alignment = align_center
        ws1.row_dimensions[1].height = 28

        ws1.merge_cells("A2:I2")
        ws1["A2"].value = "Temporary passwords must be changed on first login. Do not share publicly or commit this file to source control."
        ws1["A2"].font = font_notice
        ws1["A2"].fill = fill_notice
        ws1["A2"].alignment = align_center
        ws1.row_dimensions[2].height = 20

        ws1.row_dimensions[3].height = 10

        headers_1 = [
            "S.No.", "Role", "Class", "Register Number / Employee ID",
            "Username / Login ID", "Full Name", "Temporary Password",
            "Force Password Change", "Account Status"
        ]
        ws1.append([])
        ws1.append(headers_1)
        ws1.row_dimensions[4].height = 26

        for col_idx in range(1, len(headers_1) + 1):
            cell = ws1.cell(row=4, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        all_accounts = faculty_cred_list + student_cred_list
        for idx, item in enumerate(all_accounts, 1):
            row_num = idx + 4
            row_data = [
                idx,
                item["role"],
                item["class_name"],
                item["identifier"],
                item["username"],
                item["full_name"],
                item["temp_password"],
                item["force_password_change"],
                item["account_status"]
            ]
            ws1.append(row_data)
            ws1.row_dimensions[row_num].height = 20
            fill_row = fill_even if idx % 2 == 0 else fill_odd

            for col_idx in range(1, len(row_data) + 1):
                cell = ws1.cell(row=row_num, column=col_idx)
                cell.fill = fill_row
                cell.border = thin_border
                if col_idx in [1, 2, 4, 5, 8, 9]:
                    cell.alignment = align_center
                    cell.font = font_data
                elif col_idx == 7:  # Password
                    cell.alignment = align_center
                    cell.font = font_password
                else:
                    cell.alignment = align_left
                    cell.font = font_data

        ws1.auto_filter.ref = f"A4:I{len(all_accounts) + 4}"
        ws1.freeze_panes = "A5"

        for col in ws1.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row >= 4 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # -----------------------------
        # Sheet 2: Students
        # -----------------------------
        ws2 = wb.create_sheet(title="Students")
        ws2.views.sheetView[0].showGridLines = True

        ws2.merge_cells("A1:H1")
        ws2["A1"].value = "CONFIDENTIAL — TEMPORARY LOGIN CREDENTIALS (STUDENTS)"
        ws2["A1"].font = font_title
        ws2["A1"].fill = fill_title
        ws2["A1"].alignment = align_center
        ws2.row_dimensions[1].height = 28

        ws2.merge_cells("A2:H2")
        ws2["A2"].value = "Temporary passwords must be changed on first login. Do not share publicly or commit this file to source control."
        ws2["A2"].font = font_notice
        ws2["A2"].fill = fill_notice
        ws2["A2"].alignment = align_center
        ws2.row_dimensions[2].height = 20

        ws2.row_dimensions[3].height = 10
        headers_2 = [
            "S.No.", "Register Number", "Full Name", "Username",
            "Temporary Password", "Class", "Force Password Change", "Account Status"
        ]
        ws2.append([])
        ws2.append(headers_2)
        ws2.row_dimensions[4].height = 26

        for col_idx in range(1, len(headers_2) + 1):
            cell = ws2.cell(row=4, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        for idx, item in enumerate(student_cred_list, 1):
            row_num = idx + 4
            row_data = [
                idx,
                item["identifier"],
                item["full_name"],
                item["username"],
                item["temp_password"],
                item["class_name"],
                item["force_password_change"],
                item["account_status"]
            ]
            ws2.append(row_data)
            ws2.row_dimensions[row_num].height = 20
            fill_row = fill_even if idx % 2 == 0 else fill_odd

            for col_idx in range(1, len(row_data) + 1):
                cell = ws2.cell(row=row_num, column=col_idx)
                cell.fill = fill_row
                cell.border = thin_border
                if col_idx in [1, 2, 4, 7, 8]:
                    cell.alignment = align_center
                    cell.font = font_data
                elif col_idx == 5:
                    cell.alignment = align_center
                    cell.font = font_password
                else:
                    cell.alignment = align_left
                    cell.font = font_data

        ws2.auto_filter.ref = f"A4:H{len(student_cred_list) + 4}"
        ws2.freeze_panes = "A5"

        for col in ws2.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row >= 4 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # -----------------------------
        # Sheet 3: Faculty
        # -----------------------------
        ws3 = wb.create_sheet(title="Faculty")
        ws3.views.sheetView[0].showGridLines = True

        ws3.merge_cells("A1:H1")
        ws3["A1"].value = "CONFIDENTIAL — TEMPORARY LOGIN CREDENTIALS (FACULTY)"
        ws3["A1"].font = font_title
        ws3["A1"].fill = fill_title
        ws3["A1"].alignment = align_center
        ws3.row_dimensions[1].height = 28

        ws3.merge_cells("A2:H2")
        ws3["A2"].value = "Temporary passwords must be changed on first login. Do not share publicly or commit this file to source control."
        ws3["A2"].font = font_notice
        ws3["A2"].fill = fill_notice
        ws3["A2"].alignment = align_center
        ws3.row_dimensions[2].height = 20

        ws3.row_dimensions[3].height = 10
        headers_3 = [
            "S.No.", "Employee ID", "Full Name", "Username",
            "Temporary Password", "Department", "Force Password Change", "Account Status"
        ]
        ws3.append([])
        ws3.append(headers_3)
        ws3.row_dimensions[4].height = 26

        for col_idx in range(1, len(headers_3) + 1):
            cell = ws3.cell(row=4, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        for idx, item in enumerate(faculty_cred_list, 1):
            row_num = idx + 4
            row_data = [
                idx,
                item["identifier"],
                item["full_name"],
                item["username"],
                item["temp_password"],
                item["department"],
                item["force_password_change"],
                item["account_status"]
            ]
            ws3.append(row_data)
            ws3.row_dimensions[row_num].height = 20
            fill_row = fill_even if idx % 2 == 0 else fill_odd

            for col_idx in range(1, len(row_data) + 1):
                cell = ws3.cell(row=row_num, column=col_idx)
                cell.fill = fill_row
                cell.border = thin_border
                if col_idx in [1, 2, 4, 7, 8]:
                    cell.alignment = align_center
                    cell.font = font_data
                elif col_idx == 5:
                    cell.alignment = align_center
                    cell.font = font_password
                else:
                    cell.alignment = align_left
                    cell.font = font_data

        ws3.auto_filter.ref = f"A4:H{len(faculty_cred_list) + 4}"
        ws3.freeze_panes = "A5"

        for col in ws3.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row >= 4 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws3.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.remove(default_sheet)
        return wb

    # Save to canonical path and timestamped path
    wb1 = build_workbook()
    wb1.save(canonical_excel_path)
    print(f"  Canonical Excel saved to: {canonical_excel_path}")

    wb2 = build_workbook()
    wb2.save(timestamped_excel_path)
    print(f"  Timestamped Excel saved to: {timestamped_excel_path}")

    # -------------------------------------------------------------
    # 6. EXCEL PROGRAMMATIC VALIDATION
    # -------------------------------------------------------------
    print("\n[Step 6] Validating Excel File from Disk...")
    wb_verify = openpyxl.load_workbook(canonical_excel_path)
    assert wb_verify.sheetnames == ["All Credentials", "Students", "Faculty"]

    ws_all = wb_verify["All Credentials"]
    all_data_rows = list(ws_all.iter_rows(values_only=True))[4:]
    assert len(all_data_rows) == 127, f"Expected 127 rows in All Credentials, got {len(all_data_rows)}"

    ws_stu = wb_verify["Students"]
    stu_data_rows = list(ws_stu.iter_rows(values_only=True))[4:]
    assert len(stu_data_rows) == 119, f"Expected 119 rows in Students, got {len(stu_data_rows)}"

    ws_fac = wb_verify["Faculty"]
    fac_data_rows = list(ws_fac.iter_rows(values_only=True))[4:]
    assert len(fac_data_rows) == 8, f"Expected 8 rows in Faculty, got {len(fac_data_rows)}"

    seen_excel_unames = set()
    seen_excel_pwds = set()
    for r in all_data_rows:
        uname = r[4]
        pwd = r[6]
        fpc = r[7]
        stat = r[8]
        assert uname not in seen_excel_unames, f"Duplicate username in Excel: {uname}"
        seen_excel_unames.add(uname)
        assert pwd not in seen_excel_pwds, "Duplicate password in Excel!"
        seen_excel_pwds.add(pwd)
        assert len(pwd) == 8
        assert fpc == "TRUE"
        assert stat == "ACTIVE"

    print("  Excel Programmatic Validation: PASS (3 sheets, 127 total rows, 119 students, 8 faculty, 127 unique credentials)")

    # -------------------------------------------------------------
    # 7. 100% FULL-SYSTEM BACKEND AUTHENTICATION VERIFICATION
    # -------------------------------------------------------------
    print("\n[Step 7] 100% Backend Authentication Verification (Mandatory for EVERY Account)...")
    print("  Note: Testing against live FastAPI backend at " + BACKEND_BASE_URL)
    
    verified_faculty = 0
    verified_students = 0
    mismatched_accounts = []

    for r_idx, row in enumerate(all_data_rows, 1):
        s_no = row[0]
        role = row[1]
        cname = row[2]
        ident = str(row[3])
        uname = str(row[4])
        fname = str(row[5])
        excel_pwd = str(row[6])
        excel_fpc = row[7]
        excel_stat = row[8]

        if role == "FACULTY":
            # 1. Login via /api/v1/auth/login
            status, login_res = http_post(
                f"{BACKEND_BASE_URL}/api/v1/auth/login",
                {"username_or_email": uname, "password": excel_pwd}
            )
            if status != 200:
                mismatched_accounts.append((s_no, uname, role, f"Login failed HTTP {status}: {login_res}"))
                continue

            access_token = login_res.get("access_token")
            if not access_token:
                mismatched_accounts.append((s_no, uname, role, "Missing access_token in login response"))
                continue

            # 2. Verify identity via /api/v1/auth/me
            me_status, me_res = http_get(f"{BACKEND_BASE_URL}/api/v1/auth/me", token=access_token)
            if me_status != 200:
                mismatched_accounts.append((s_no, uname, role, f"/auth/me failed HTTP {me_status}: {me_res}"))
                continue

            # 3. Confirm identity matches expected user
            if me_res.get("role") != "FACULTY":
                mismatched_accounts.append((s_no, uname, role, f"Role mismatch: {me_res.get('role')} != FACULTY"))
                continue
            if me_res.get("username").lower() != uname.lower():
                mismatched_accounts.append((s_no, uname, role, f"Username mismatch: {me_res.get('username')} != {uname}"))
                continue
            if not me_res.get("is_active"):
                mismatched_accounts.append((s_no, uname, role, "Account not active"))
                continue

            # 4. Confirm DB force_password_change
            cur.execute("SELECT force_password_change, is_active FROM users WHERE username = ?", (uname,))
            db_user_row = cur.fetchone()
            if not db_user_row or db_user_row[0] != 1 or db_user_row[1] != 1:
                mismatched_accounts.append((s_no, uname, role, f"DB state mismatch: {db_user_row}"))
                continue

            verified_faculty += 1
            if verified_faculty % 2 == 0 or verified_faculty == 8:
                print(f"    [Faculty {verified_faculty}/8] Verified: {ident} ({uname}) -> HTTP 200, Role: FACULTY, Active: True, FPC: True")

        elif role == "STUDENT":
            # 1. Login via /api/v1/auth/student/login
            status, login_res = http_post(
                f"{BACKEND_BASE_URL}/api/v1/auth/student/login",
                {"identifier": ident, "password": excel_pwd}
            )
            if status != 200:
                mismatched_accounts.append((s_no, uname, role, f"Student login failed HTTP {status}: {login_res}"))
                continue

            access_token = login_res.get("access_token")
            returned_fpc = login_res.get("force_password_change")
            returned_reg = login_res.get("register_number")

            if not access_token:
                mismatched_accounts.append((s_no, uname, role, "Missing access_token in login response"))
                continue
            if returned_fpc is not True:
                mismatched_accounts.append((s_no, uname, role, f"force_password_change in response is {returned_fpc}"))
                continue
            if returned_reg != ident:
                mismatched_accounts.append((s_no, uname, role, f"Register number mismatch: {returned_reg} != {ident}"))
                continue

            # 2. Verify identity via /api/v1/auth/me
            me_status, me_res = http_get(f"{BACKEND_BASE_URL}/api/v1/auth/me", token=access_token)
            if me_status != 200:
                mismatched_accounts.append((s_no, uname, role, f"/auth/me failed HTTP {me_status}: {me_res}"))
                continue

            if me_res.get("role") != "STUDENT":
                mismatched_accounts.append((s_no, uname, role, f"Role mismatch: {me_res.get('role')} != STUDENT"))
                continue
            if me_res.get("username").lower() != uname.lower():
                mismatched_accounts.append((s_no, uname, role, f"Username mismatch: {me_res.get('username')} != {uname}"))
                continue
            if not me_res.get("is_active"):
                mismatched_accounts.append((s_no, uname, role, "Account not active"))
                continue

            # 3. Confirm DB force_password_change
            cur.execute("SELECT force_password_change, is_active FROM users WHERE username = ?", (uname,))
            db_user_row = cur.fetchone()
            if not db_user_row or db_user_row[0] != 1 or db_user_row[1] != 1:
                mismatched_accounts.append((s_no, uname, role, f"DB state mismatch: {db_user_row}"))
                continue

            verified_students += 1
            if verified_students % 25 == 0 or verified_students == 119:
                print(f"    [Student {verified_students}/119] Verified: {ident} ({uname}) -> HTTP 200, Role: STUDENT, Active: True, FPC: True")

    total_verified = verified_faculty + verified_students
    print(f"\n  Backend Authentication Verification Results:")
    print(f"    - Faculty Verified: {verified_faculty}/8")
    print(f"    - Students Verified: {verified_students}/119")
    print(f"    - Total Accounts Verified: {total_verified}/127")

    if mismatched_accounts:
        print(f"  FAILED ACCOUNTS ({len(mismatched_accounts)}):")
        for err in mismatched_accounts[:10]:
            print(f"    {err}")
        raise RuntimeError(f"Backend authentication failed for {len(mismatched_accounts)} accounts!")

    assert verified_faculty == 8
    assert verified_students == 119
    assert total_verified == 127
    print("  100% Backend Authentication Verification: PASS (127/127 Accounts Successfully Authenticated)")

    # -------------------------------------------------------------
    # 8. SECURITY & RBAC ISOLATION TESTING
    # -------------------------------------------------------------
    print("\n[Step 8] Testing Security & RBAC Isolation...")
    
    # 1. Admin login verification (must succeed with existing credentials)
    adm_status, adm_res = http_post(
        f"{BACKEND_BASE_URL}/api/v1/auth/login",
        {"username_or_email": "admin", "password": "admin123"}
    )
    if adm_status != 200:
        raise RuntimeError(f"Admin login failed: HTTP {adm_status} {adm_res}")
    adm_token = adm_res.get("access_token")
    adm_me_st, adm_me = http_get(f"{BACKEND_BASE_URL}/api/v1/auth/me", token=adm_token)
    assert adm_me_st == 200 and adm_me.get("role") == "ADMIN"
    print("  Admin Login & RBAC: PASS (HTTP 200, Role: ADMIN)")

    # 2. Invalid password rejection
    sample_fac = faculty_cred_list[0]
    inv_f_st, inv_f_res = http_post(
        f"{BACKEND_BASE_URL}/api/v1/auth/login",
        {"username_or_email": sample_fac["username"], "password": "WrongPassword999!"}
    )
    assert inv_f_st == 401, f"Expected 401 for wrong faculty password, got {inv_f_st}"
    print("  Invalid Faculty Password Rejection: PASS (HTTP 401)")

    sample_stu = student_cred_list[0]
    inv_s_st, inv_s_res = http_post(
        f"{BACKEND_BASE_URL}/api/v1/auth/student/login",
        {"identifier": sample_stu["identifier"], "password": "WrongPassword999!"}
    )
    assert inv_s_st == 401, f"Expected 401 for wrong student password, got {inv_s_st}"
    print("  Invalid Student Password Rejection: PASS (HTTP 401)")

    # 3. RBAC route protection
    # Get student token
    s_st, s_res = http_post(
        f"{BACKEND_BASE_URL}/api/v1/auth/student/login",
        {"identifier": sample_stu["identifier"], "password": sample_stu["temp_password"]}
    )
    s_token = s_res["access_token"]
    rbac_s_st, _ = http_get(f"{BACKEND_BASE_URL}/api/v1/faculty/timetable", token=s_token)
    assert rbac_s_st == 403, f"Expected 403 for student accessing faculty route, got {rbac_s_st}"
    print("  Student RBAC Route Isolation: PASS (HTTP 403 on faculty timetable)")

    # Get faculty token
    f_st, f_res = http_post(
        f"{BACKEND_BASE_URL}/api/v1/auth/login",
        {"username_or_email": sample_fac["username"], "password": sample_fac["temp_password"]}
    )
    f_token = f_res["access_token"]
    rbac_f_st, _ = http_get(f"{BACKEND_BASE_URL}/api/v1/admin/students", token=f_token)
    assert rbac_f_st == 403, f"Expected 403 for faculty accessing admin route, got {rbac_f_st}"
    print("  Faculty RBAC Route Isolation: PASS (HTTP 403 on admin endpoint)")

    # -------------------------------------------------------------
    # 9. POST-OPERATION DATABASE INVARIANTS AUDIT
    # -------------------------------------------------------------
    print("\n[Step 9] Post-Operation Database Invariants Audit...")
    
    # Clean up test session tokens so all accounts are in logged-out state
    cur.execute("UPDATE users SET session_token = NULL, failed_login_attempts = 0, locked_until = NULL")
    conn.commit()

    cur.execute("SELECT count(*) FROM students")
    post_students = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM faculty")
    post_faculty = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM users")
    post_users = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM class_sections")
    post_sections = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM subjects")
    post_subjects = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM attendance_records")
    post_attendance = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM class_sessions")
    post_sessions = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM face_embeddings")
    post_embeddings = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM users WHERE force_password_change = 1")
    post_fpc = cur.fetchone()[0]

    cur.execute("SELECT count(*) FROM users WHERE role = 'ADMIN' AND force_password_change = 1")
    post_admin_fpc = cur.fetchone()[0]

    print(f"  Students Total: {post_students} (Expected: 119, Delta: {post_students - pre_students})")
    print(f"  Faculty Total: {post_faculty} (Expected: 8, Delta: {post_faculty - pre_faculty})")
    print(f"  Users Total: {post_users} (Expected: 128, Delta: {post_users - pre_users})")
    print(f"  Class Sections: {post_sections} (Expected: 2, Delta: 0)")
    print(f"  Subjects: {post_subjects} (Expected: 27, Delta: 0)")
    print(f"  Attendance Records: {post_attendance} (Expected: 0, Delta: 0)")
    print(f"  Class Sessions: {post_sessions} (Expected: 0, Delta: 0)")
    print(f"  Face Embeddings: {post_embeddings} (Expected: 0, Delta: 0)")
    print(f"  Total Accounts with force_password_change=1: {post_fpc} (Expected: 127)")
    print(f"  Admin force_password_change=1: {post_admin_fpc} (Expected: 0)")

    assert post_students == 119
    assert post_faculty == 8
    assert post_users == 128
    assert post_sections == 2
    assert post_subjects == 27
    assert post_attendance == 0
    assert post_sessions == 0
    assert post_embeddings == 0
    assert post_fpc == 127
    assert post_admin_fpc == 0

    # Verify no plaintext passwords in DB
    cur.execute("SELECT hashed_password FROM users")
    all_hashes = [r[0] for r in cur.fetchall()]
    for h in all_hashes:
        assert h.startswith("$2b$") or h.startswith("$2a$"), "Found non-bcrypt password in DB!"
    print("  Password Storage Security: PASS (All passwords stored as bcrypt hashes, 0 plaintext passwords)")

    conn.close()

    print("\n" + "=" * 70)
    print("SUCCESS: CREDENTIAL REGENERATION & BACKEND VERIFICATION COMPLETE")
    print(f"Excel Location: {canonical_excel_path}")
    print(f"Backup Location: {timestamped_excel_path}")
    print(f"Total Accounts Processed: {total_verified}/127")
    print("=" * 70)


if __name__ == "__main__":
    main()
