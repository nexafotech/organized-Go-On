"""Excel (.xlsx) renderer for the date-range attendance report (Phase 9).

Renders an already-built ``AttendanceRangeReport``. It performs NO attendance
calculation and runs NO database queries.
"""
from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.attendance_report_service import AttendanceRangeReport, status_label
from app.services.report_branding import BRAND_NAME, badge_path

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Summary sheet: Go-On brand block occupies rows 1-6; the student table header is row 7.
SUMMARY_HEADER_ROW = 7

_HEADER_FILL = PatternFill("solid", fgColor="4E342E")   # espresso (brand)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="BFB3A5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_DATE_FMT = "DD-MM-YYYY"
_PCT_FMT = "0.00%"


def _set_text(cell, value: str) -> None:
    """Write user-controlled text as a literal string (never evaluated as a formula)."""
    cell.value = value
    cell.data_type = "s"
    cell.number_format = "@"


def _write_header(ws, headers, row: int = 1) -> None:
    for col, title in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=title)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.border = _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = f"A{row + 1}"


def _set_widths(ws, widths) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _print_setup(ws, report: AttendanceRangeReport, landscape: bool = False, header_row: int = 1) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddHeader.center.text = (
        f"Attendance Report - {report.class_name} - {report.subject_label} "
        f"({report.from_date:%d-%m-%Y} to {report.to_date:%d-%m-%Y})"
    )
    ws.oddFooter.center.text = "Page &P of &N"


def _write_brand_block(ws, report: AttendanceRangeReport) -> None:
    """Go-On logo + name + report title + Class / Subject / Date Range (rows 1-5)."""
    espresso = Font(bold=True, size=18, color="4E342E")
    ws.row_dimensions[1].height = 27
    ws.row_dimensions[2].height = 21
    ws.row_dimensions[6].height = 8

    _set_text(ws.cell(row=1, column=2), BRAND_NAME)
    ws.cell(row=1, column=2).font = espresso
    ws.cell(row=1, column=2).alignment = Alignment(vertical="bottom")
    _set_text(ws.cell(row=2, column=2), "Attendance Report")
    ws.cell(row=2, column=2).font = Font(bold=True, size=12, color="6D5D52")
    ws.cell(row=2, column=2).alignment = Alignment(vertical="top")

    meta = [
        ("Class", report.class_name),
        ("Subject", report.subject_label),
        ("Date Range", f"{report.from_date:%d-%m-%Y} to {report.to_date:%d-%m-%Y}"),
    ]
    for offset, (label, value) in enumerate(meta, start=3):
        ws.cell(row=offset, column=1, value=label).font = Font(bold=True)
        _set_text(ws.cell(row=offset, column=2), value)

    logo = badge_path()
    if logo is not None:
        img = XLImage(str(logo))
        img.width = img.height = 46          # keep square; source is 256px (light file)
        ws.add_image(img, "A1")


def render_excel(report: AttendanceRangeReport) -> bytes:
    wb = Workbook()

    # ── Sheet 1: Summary ────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"

    # Go-On brand block (rows 1-5), then a spacer; the table header sits on row 7.
    _write_brand_block(ws, report)
    headers = ["Register Number", "Student Name", "Present", "Absent", "OD", "Total Marked", "Attendance %"]
    _write_header(ws, headers, row=SUMMARY_HEADER_ROW)
    r = SUMMARY_HEADER_ROW
    for row in report.summary:
        r += 1
        _set_text(ws.cell(row=r, column=1), row.roll_number)
        _set_text(ws.cell(row=r, column=2), row.full_name)
        ws.cell(row=r, column=3, value=row.present)
        ws.cell(row=r, column=4, value=row.absent)
        ws.cell(row=r, column=5, value=row.od)
        ws.cell(row=r, column=6, value=row.total_marked)
        pc = ws.cell(row=r, column=7)
        if row.percentage is None:
            pc.value = "N/A"   # no marked entries: NOT 0%
            pc.alignment = Alignment(horizontal="right")
        else:
            pc.value = round(row.percentage / 100.0, 6)
            pc.number_format = _PCT_FMT
        for col in range(1, 8):
            ws.cell(row=r, column=col).border = _BORDER
    last_data_row = r
    ws.auto_filter.ref = f"A{SUMMARY_HEADER_ROW}:G{max(last_data_row, SUMMARY_HEADER_ROW)}"

    # Totals + context block, outside the filter range (one blank row between)
    r = last_data_row + 2
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=r, column=3, value=report.present).font = Font(bold=True)
    ws.cell(row=r, column=4, value=report.absent).font = Font(bold=True)
    ws.cell(row=r, column=5, value=report.od).font = Font(bold=True)
    ws.cell(row=r, column=6, value=report.total_marked).font = Font(bold=True)
    tp = ws.cell(row=r, column=7)
    if report.percentage is None:
        tp.value = "N/A"
        tp.alignment = Alignment(horizontal="right")
    else:
        tp.value = round(report.percentage / 100.0, 6)
        tp.number_format = _PCT_FMT
    tp.font = Font(bold=True)

    info = [
        ("Generated", report.generated_at.strftime("%d-%m-%Y %H:%M")),
        ("Percentage basis", report.percentage_basis),
    ]
    if report.scope_note:
        info.append(("Scope", report.scope_note))
    r += 2
    for label, value in info:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        _set_text(ws.cell(row=r, column=2), value)
        r += 1
    _set_widths(ws, [20, 34, 10, 10, 8, 14, 15])
    _print_setup(ws, report, header_row=SUMMARY_HEADER_ROW)

    # ── Sheet 2: Daily Attendance ───────────────────────────
    wd = wb.create_sheet("Daily Attendance")
    _write_header(wd, ["Date", "Period", "Subject", "Register Number", "Student Name", "Status"])
    r = 1
    for d in report.details:
        r += 1
        dc = wd.cell(row=r, column=1, value=d.date)
        dc.number_format = _DATE_FMT
        dc.alignment = Alignment(horizontal="left")
        _set_text(wd.cell(row=r, column=2), f"P{d.period_number}")
        _set_text(wd.cell(row=r, column=3), d.subject_code)
        _set_text(wd.cell(row=r, column=4), d.roll_number)
        _set_text(wd.cell(row=r, column=5), d.full_name)
        _set_text(wd.cell(row=r, column=6), status_label(d.status))
        for col in range(1, 7):
            wd.cell(row=r, column=col).border = _BORDER
    wd.auto_filter.ref = f"A1:F{max(r, 1)}"
    _set_widths(wd, [13, 9, 14, 20, 34, 16])
    _print_setup(wd, report)

    # ── Sheet 3: Period Summary ─────────────────────────────
    wp = wb.create_sheet("Period Summary")
    _write_header(wp, ["Date", "Period", "Subject", "Present", "Absent", "OD", "Total Marked"])
    r = 1
    for p in report.period_summary:
        r += 1
        dc = wp.cell(row=r, column=1, value=p.date)
        dc.number_format = _DATE_FMT
        dc.alignment = Alignment(horizontal="left")
        _set_text(wp.cell(row=r, column=2), f"P{p.period_number}")
        _set_text(wp.cell(row=r, column=3), p.subject_code)
        wp.cell(row=r, column=4, value=p.present)
        wp.cell(row=r, column=5, value=p.absent)
        wp.cell(row=r, column=6, value=p.od)
        wp.cell(row=r, column=7, value=p.total_marked)
        for col in range(1, 8):
            wp.cell(row=r, column=col).border = _BORDER
    wp.auto_filter.ref = f"A1:G{max(r, 1)}"
    _set_widths(wp, [13, 9, 14, 10, 10, 8, 14])
    _print_setup(wp, report)

    wb.properties.title = f"Attendance Report - {report.class_name}"
    wb.properties.subject = report.subject_label
    wb.properties.creator = report.institution

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
